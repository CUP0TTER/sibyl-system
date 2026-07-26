from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path('generated/建筑设计基础2')
OUT.mkdir(parents=True, exist_ok=True)
OUTPUT = OUT / '2023级建筑学_建筑设计基础2_教学进度表_16周96学时_严格模板版.docx'

COURSE = '建筑设计基础2'
CLASS_NAME = '2023级建筑学'
TEACHERS = '杨巧梅、周枳辛'
ROOM = '1B-305'
START = date(2024, 2, 26)  # 第一周周一

# 严格依据大纲16周安排：每周周一3学时、周三3学时，共6学时；16周共96学时。
THEORY = [
    '建筑抄绘理论与案例剖析：建筑抄绘的目标、方法、关键要点及经典建筑案例分析',
    '建筑抄绘实践（一）：建筑平面图识读，比例、尺度、轴线、功能与空间组织分析',
    '建筑抄绘实践（一）：建筑立面、剖面及透视表达，形体、构造与空间关系分析',
    '建筑抄绘实践（二）：空间序列、功能布局、流线组织、比例控制与图线层级深化',
    '建筑抄绘实践（二）：成果完善、建筑分析、设计元素提炼及抄绘总结方法',
    '平面构成原理与练习（一）：平面构成概念，点、线、面基本元素及视觉特征',
    '平面构成原理与练习（二）：对称、均衡、重复、渐变、发射等形式法则',
    '平面构成原理与练习（三）：对比与调和、节奏韵律、比例、正负形及共生图形',
    '立体构成原理与练习（一）：立体构成概念、块体组合、空间限定及体量关系',
    '立体构成原理与练习（二）：线材、面材构成，材料特性、连接方式与结构稳定',
    '立体构成原理与练习（三）：综合形态、空间序列、比例、光影及作品表达评析',
    '模型制作材料与工具讲解：模型比例、材料性能、工具使用、安全操作及制作流程',
    '模型制作实践（一）：设计图纸识读、模型分解、底板定位、主体结构与墙体制作',
    '模型制作实践（一）：楼板、屋面、门窗、立面构件及主要空间关系的模型表达',
    '模型制作实践（二）：细部完善、材质色彩、场地环境、景观构件及整体效果控制',
    '模型制作实践（二）与课程总结：成果展示、设计说明、模型摄影、评价与学习复盘',
]

PRACTICE = [
    '建筑抄绘准备：熟悉绘图工具与线型，选择案例，确定比例、图幅和版面',
    '完成经典建筑平面图抄绘，校核比例、尺度、轴线、功能和空间关系',
    '完成建筑立面图、剖面图及外观透视图抄绘，统一图线和标注表达',
    '深化平、立、剖图纸，修正比例、构造细节、线条层级和图面问题',
    '完成建筑分析图、抄绘心得和成果版面，开展阶段展示与集中讲评',
    '开展点、线、面组合及基本骨格练习，形成多组平面构成草案',
    '运用对称、均衡、重复、渐变等法则完成主题平面构成练习',
    '完成正负形、共生图形或综合构成，优化版面并进行成果讲评',
    '完成块体切割、组合与空间限定练习，比较不同体量和空间关系',
    '完成线材、面材立体构成，重点检查节点连接和结构稳定性',
    '完成综合立体构成作品，调整空间、比例、光影和展示方式',
    '开展材料切割、粘接与工具安全训练，确定模型比例和制作计划',
    '制作模型底板、定位线、主体结构、墙体及主要空间构件',
    '完成楼板、屋面、门窗和立面构件，校核比例与空间关系',
    '完善材质、色彩、场地及景观细部，修正模型制作工艺问题',
    '完成最终模型、成果摄影、展示汇报、评价反馈和成果归档',
]

assert len(THEORY) == 16 and len(PRACTICE) == 16


def set_font(run, size=8.0, bold=False, name='宋体'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold


def set_cell_text(cell, text, size=7.2, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(str(text))
    set_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_margins(cell, value=25):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side in ('top', 'start', 'bottom', 'end'):
        node = tc_mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    tc_w.set(qn('w:w'), str(int(width_cm * 567)))
    tc_w.set(qn('w:type'), 'dxa')


def fixed_layout(table):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        tbl_pr.append(layout)
    layout.set(qn('w:type'), 'fixed')


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement('w:cantSplit')
    tr_pr.append(tag)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        node = borders.find(qn(f'w:{edge}'))
        if node is None:
            node = OxmlElement(f'w:{edge}')
            borders.append(node)
        node.set(qn('w:val'), 'nil')


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(0.7)
    section.bottom_margin = Cm(0.65)
    section.left_margin = Cm(0.7)
    section.right_margin = Cm(0.7)

    normal = doc.styles['Normal']
    normal.font.name = '宋体'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.font.size = Pt(8)

    # 完全沿用模板顶部结构：左侧标题，右侧学时统计表。
    top = doc.add_table(rows=1, cols=2)
    top.alignment = WD_TABLE_ALIGNMENT.CENTER
    top.autofit = False
    remove_table_borders(top)
    set_cell_width(top.cell(0, 0), 23.1)
    set_cell_width(top.cell(0, 1), 4.4)

    p = top.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run('附件 3：城乡规划教研室\n教学进度表')
    set_font(r, size=15, bold=True, name='黑体')

    summary_cell = top.cell(0, 1)
    summary_cell.text = ''
    summary = summary_cell.add_table(rows=3, cols=2)
    summary.style = 'Table Grid'
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary.autofit = False
    for row in summary.rows:
        row.height = Cm(0.62)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    for i, (left, right) in enumerate([
        ('计划总学时', '96学时'),
        ('理论课学时', '48学时'),
        ('实验课学时', '48学时'),
    ]):
        set_cell_width(summary.cell(i, 0), 2.2)
        set_cell_width(summary.cell(i, 1), 2.2)
        set_cell_text(summary.cell(i, 0), left, size=8.2)
        set_cell_text(summary.cell(i, 1), right, size=8.2)

    info = doc.add_paragraph()
    info.paragraph_format.space_before = Pt(2)
    info.paragraph_format.space_after = Pt(2)
    r = info.add_run('课程名称 建筑设计基础2  班级 2023级建筑学  从 2024 年 2 月 26 日至 2024 年 6 月 14 日止')
    set_font(r, size=9.2)

    # 主表严格保持模板17列、表格线、表头文字和左右分栏结构。
    widths = [1.52, 0.95, 0.60, 0.60, 0.45, 0.45, 0.56, 7.48, 0.72,
              1.52, 0.95, 0.60, 0.45, 0.45, 0.56, 7.05, 0.78]
    table = doc.add_table(rows=2, cols=17)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    fixed_layout(table)

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
            set_cell_margins(cell)

    table.cell(0, 0).merge(table.cell(0, 8))
    set_cell_text(table.cell(0, 0), '理论课教学安排', size=9.0, bold=True)
    table.cell(0, 9).merge(table.cell(0, 16))
    set_cell_text(table.cell(0, 9), '实验课教学安排', size=9.0, bold=True)

    headers = [
        '教师姓名', '职称', '周次', '年', '月', '日', '时数', '讲课章节及内容', '备注',
        '教师姓名', '职称', '周次', '月', '日', '时数', '实验、实践内容', '学生分组',
    ]
    for i, header in enumerate(headers):
        set_cell_text(table.cell(1, i), header, size=7.4, bold=False)

    table.rows[0].height = Cm(0.62)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    table.rows[1].height = Cm(0.92)
    table.rows[1].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

    # 只设置16个教学周，每周同时填写周一和周三，共16×6=96学时。
    for week in range(1, 17):
        monday = START + timedelta(weeks=week - 1)
        wednesday = monday + timedelta(days=2)
        row = table.add_row()
        row.height = Cm(1.28)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        prevent_row_split(row)
        cells = row.cells
        for i, cell in enumerate(cells):
            set_cell_width(cell, widths[i])
            set_cell_margins(cell)

        values = [
            TEACHERS, '', week, monday.year, monday.month, monday.day, 3, THEORY[week - 1], '',
            TEACHERS, '', week, wednesday.month, wednesday.day, 3, PRACTICE[week - 1], '',
        ]
        for i, value in enumerate(values):
            left_align = i in (7, 15)
            set_cell_text(
                cells[i], value,
                size=6.5 if left_align else 7.0,
                align=WD_ALIGN_PARAGRAPH.LEFT if left_align else WD_ALIGN_PARAGRAPH.CENTER,
            )

    footer = doc.add_paragraph()
    footer.paragraph_format.space_before = Pt(2)
    footer.paragraph_format.space_after = Pt(0)
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = footer.add_run('上课时间/地点：周一下午 5、6、7 节，周三下午 7、8、9 节/1B-305')
    set_font(r, size=9.0)

    doc.save(OUTPUT)
    print(OUTPUT, OUTPUT.stat().st_size)


if __name__ == '__main__':
    build_document()
